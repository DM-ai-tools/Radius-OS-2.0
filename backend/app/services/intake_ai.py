"""AI intake fill + CDD document extraction (Perplexity Pro via OpenRouter)."""

from __future__ import annotations

import io
import re
from typing import Any

from app.config import get_settings
from app.integrations.llm import _openrouter_chat, synthesize_json
from app.integrations.web_fetch import fetch_url, page_text_excerpt, parse_html
from app.logging_config import get_logger
from app.services.discovery_fields import (
    ACCESS_FIELDS,
    CDD_LABEL_ALIASES,
    CLIENT_ONLY_FIELDS,
    RESEARCH_FIELDS,
    cdd_json_schema_hint,
)

log = get_logger("intake_ai")

# Never AI-fill personal / account-manager fields
PERSONAL_KEYS = {
    "primary_contact_name",
    "primary_contact_email",
    "primary_contact_phone",
    "account_manager",
}


async def _perplexity_json(system: str, user: str) -> dict[str, Any] | None:
    settings = get_settings()
    model = settings.research_model or "perplexity/sonar-pro"
    if settings.use_mock_llm or not settings.openrouter_api_key:
        return None
    try:
        text = await _openrouter_chat(
            system=system + "\nReturn ONLY valid JSON. No markdown fences.",
            user=user,
            model=model,
            max_tokens=4000,
        )
        if not text:
            return None
        return _parse_json_loose(text)
    except Exception as exc:  # noqa: BLE001
        log.warning("perplexity_json_failed", error=str(exc), model=model)
        # Fallback to default skill model JSON helper
        return await synthesize_json(system, user)


def _parse_json_loose(text: str) -> dict[str, Any] | None:
    import json

    raw = text.strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
    try:
        data = json.loads(raw)
        return data if isinstance(data, dict) else None
    except json.JSONDecodeError:
        m = re.search(r"\{[\s\S]*\}", raw)
        if not m:
            return None
        try:
            data = json.loads(m.group(0))
            return data if isinstance(data, dict) else None
        except json.JSONDecodeError:
            return None


async def fill_client_form_from_url(url: str) -> dict[str, Any]:
    """Research a website and return New Client form fields (no personal contacts)."""
    url = (url or "").strip()
    if not url.startswith("http"):
        url = f"https://{url}"

    excerpt = ""
    title = ""
    try:
        fetched = await fetch_url(url)
        parser = parse_html(fetched.get("text") or "")
        excerpt = page_text_excerpt(parser, 4000)
        title = parser.title or ""
    except Exception as exc:  # noqa: BLE001
        log.warning("fill_fetch_failed", error=str(exc), url=url)

    system = (
        "You are researching a company for SEO onboarding intake. "
        "Use the website URL and any public knowledge to fill business fields. "
        "Do NOT invent personal contact names, emails, or phone numbers — leave those null. "
        "Return JSON only."
    )
    user = (
        f"Website: {url}\nPage title: {title}\n"
        f"Site excerpt:\n{excerpt or '(empty — use public knowledge about this domain)'}\n\n"
        "Return JSON with keys:\n"
        "name, industry, company_size, headquarters_location, target_audience,\n"
        "geographic_focus, known_competitors, website_priorities, objectives,\n"
        "business_keywords, b2b_b2c, business_model.\n"
        "objectives = string array of SEO/marketing goals.\n"
        "known_competitors = short text with competitor names/URLs.\n"
        "Omit or null any field you cannot support."
    )
    parsed = await _perplexity_json(system, user)
    if not parsed:
        out: dict[str, Any] = {
            "primary_url": url,
            "status": "unavailable",
            "needs_review": True,
            "unavailable_reason": (
                "Discovery research could not be parsed. Needs review. "
                "No priorities, objectives, or positioning were inferred."
            ),
        }
        if title:
            observed = title.split("|")[0].split("-")[0].strip()[:80]
            if observed:
                out["name"] = observed
        return out

    out = {}
    for k, v in parsed.items():
        if k.startswith("_") or k in PERSONAL_KEYS:
            continue
        if v in (None, "", [], {}):
            continue
        out[k] = v
    if "name" not in out and title:
        out["name"] = title.split("|")[0].split("-")[0].strip()[:80]
    out["primary_url"] = url
    return out


_LLM_KEY_ALIASES = {
    "company_name": "business_name",
    "company": "business_name",
    "client_name": "business_name",
    "url": "website_url",
    "website": "website_url",
    "primary_url": "website_url",
    "site": "website_url",
    "industry": "inferred_industry",
    "products_services": "products",
    "location": "geographic_focus",
    "geo": "geographic_focus",
    "audience": "target_demographic",
    "target_audience": "target_demographic",
    "goals": "objectives",
    "goal": "business_goal",
    "competitor": "competitors",
    "known_competitors": "competitors",
    "average_order_value": "average_ticket_size",
}


def _cdd_allowed_keys() -> set[str]:
    return set(RESEARCH_FIELDS) | set(CLIENT_ONLY_FIELDS) | set(ACCESS_FIELDS) | {
        "business_name",
        "website_url",
        "competitors",
    }


def _lookup_cdd_key(label: str) -> str | None:
    raw = " ".join(str(label or "").lower().replace("_", " ").split()).strip(" :?-")
    if not raw or raw in PERSONAL_KEYS:
        return None
    key = CDD_LABEL_ALIASES.get(raw) or CDD_LABEL_ALIASES.get(raw.rstrip("?").strip())
    if not key:
        snake = raw.replace(" ", "_")
        key = _LLM_KEY_ALIASES.get(snake) or _LLM_KEY_ALIASES.get(raw)
        if snake in _cdd_allowed_keys():
            key = snake
    if not key:
        # Prefer the label containing the alias (more specific wording). Do not
        # let a short token such as "business" claim every "business …" alias.
        for alias, target in CDD_LABEL_ALIASES.items():
            if len(alias) < 8:
                continue
            if alias == raw or alias in raw:
                key = target
                break
            if len(raw) >= 12 and raw in alias:
                key = target
                break
    if key == "industry":
        key = "inferred_industry"
    if key in PERSONAL_KEYS or key not in _cdd_allowed_keys():
        return None
    return key


def _coerce_cdd_payload(parsed: Any) -> dict[str, Any]:
    """Flatten model JSON into a single field map, including nested or list shapes."""
    if isinstance(parsed, list):
        flat: dict[str, Any] = {}
        for item in parsed:
            flat.update(_coerce_cdd_payload(item))
        return flat
    if not isinstance(parsed, dict):
        return {}
    if any(k in parsed for k in ("field", "key", "label")) and "value" in parsed:
        label = parsed.get("field") or parsed.get("key") or parsed.get("label")
        return {str(label): parsed.get("value")}
    nested_keys = ("fields", "cdd", "data", "discovery", "answers", "result")
    if len(parsed) == 1:
        only = next(iter(parsed.values()))
        if isinstance(only, (dict, list)) and next(iter(parsed)) in nested_keys:
            return _coerce_cdd_payload(only)
    out: dict[str, Any] = {}
    for key, value in parsed.items():
        if isinstance(value, dict) and key in nested_keys:
            out.update(_coerce_cdd_payload(value))
            continue
        if isinstance(value, list) and value and all(isinstance(item, dict) for item in value):
            out.update(_coerce_cdd_payload(value))
            continue
        out[str(key)] = value
    return out


async def extract_cdd_fields_from_text(text: str, *, filename: str = "") -> dict[str, Any]:
    """Map free-text / spreadsheet dump into CDD discovery field keys."""
    system = (
        "You extract Client Discovery Document (CDD / APSA) fields for SEO onboarding. "
        "Map whatever is present into the JSON keys listed. "
        "Do not invent personal contact details. Return JSON only."
    )
    user = (
        f"Source file: {filename or 'upload'}\n"
        f"Allowed keys: {cdd_json_schema_hint()}, business_name, website_url, competitors\n\n"
        f"Document text:\n{text[:18000]}\n\n"
        "Return one JSON object keyed by those fields. Use null when unknown."
    )
    parsed = await _perplexity_json(system, user)
    from_model = _sanitize_cdd(_coerce_cdd_payload(parsed or {}))
    # Labeled cells in the file win. The model fills keys the sheet did not
    # state, and still recovers a narrative PDF when there are no labels.
    from_sheet = _sanitize_cdd(_heuristic_cdd_map(text))
    merged = dict(from_model)
    for key, value in from_sheet.items():
        if str(value or "").strip():
            merged[key] = value
    if not merged:
        log.warning(
            "cdd_extract_empty",
            filename=filename,
            text_chars=len(text or ""),
            model_keys=list((parsed or {}).keys())[:12] if isinstance(parsed, dict) else [],
        )
    return merged


def _is_document_stamp(value: str) -> bool:
    """CDD header lines such as '21/09/2026 (draft)' are not field values."""
    text = " ".join(str(value or "").split()).strip()
    if not text or text.lower() in {"draft", "cdd", "untitled", "client discovery document"}:
        return True
    return bool(
        re.match(
            r"^(?:\d{1,2}[/.\\-]\d{1,2}[/.\\-]\d{2,4})(?:\s*\((?:draft|final|v\d+)\))?$",
            text,
            re.IGNORECASE,
        )
    )


def _sanitize_cdd(parsed: dict[str, Any]) -> dict[str, Any]:
    allowed = _cdd_allowed_keys()
    out: dict[str, Any] = {}
    for raw_key, value in parsed.items():
        key = _lookup_cdd_key(str(raw_key))
        if not key or key not in allowed or key in PERSONAL_KEYS:
            continue
        if isinstance(value, dict) or value in (None, "", [], {}):
            continue
        if isinstance(value, list):
            value = ", ".join(str(item).strip() for item in value if str(item).strip())
        if not str(value).strip():
            continue
        if _is_document_stamp(str(value)):
            # Spreadsheet titles and file dates are not industry, products, or names.
            continue
        out[key] = value
    return out


def _pipe_label_value_pairs(raw: str) -> list[tuple[str, str]]:
    """Read label/value pairs from a pipe-joined spreadsheet row.

    APSA CDDs are usually two columns. Extra trailing cells (dates, notes,
    status) must not replace the value for the first label — that was how
    ``21/09/2026 (draft)`` landed in Industry, Products, and Business Name.
    """
    parts = [p.strip() for p in raw.split("|")]
    nonempty = [p for p in parts if p]
    if len(nonempty) < 2:
        return []
    pairs: list[tuple[str, str]] = []
    index = 0
    while index < len(nonempty) - 1:
        label = nonempty[index]
        value = nonempty[index + 1]
        if _lookup_cdd_key(label) and not _lookup_cdd_key(value):
            pairs.append((label, value))
            index += 2
            continue
        index += 1
    return pairs


def _heuristic_cdd_map(text: str) -> dict[str, Any]:
    out: dict[str, Any] = {}
    lines = [line.strip() for line in text.splitlines()]
    index = 0
    access_keys = {
        "analytics_access",
        "gtm_access",
        "google_ads_access",
        "merchant_center_access",
        "access_level",
        "search_console_access",
        "cms_access",
        "hosting_access",
        "google_business_access",
    }
    while index < len(lines):
        raw = lines[index]
        index += 1
        if not raw or raw.startswith("##"):
            continue

        row_pairs: list[tuple[str, str]] = []
        if "|" in raw:
            row_pairs = _pipe_label_value_pairs(raw)
            if not row_pairs:
                continue
        elif ":" in raw:
            label, value = raw.split(":", 1)
            row_pairs = [(label.strip(), value.strip())]
        else:
            row_pairs = [(raw, "")]

        for label, value in row_pairs:
            if not value and "(" in label and ")" in label:
                start = label.rfind("(")
                end = label.rfind(")")
                if start >= 0 and end > start:
                    value = label[start + 1 : end].strip()
                    label = label[:start].strip()

            key = _lookup_cdd_key(label)
            if key and not value and index < len(lines) and len(row_pairs) == 1:
                nxt = lines[index]
                if nxt and not nxt.startswith("##") and not _lookup_cdd_key(nxt.split("|")[0]):
                    value = nxt.split("|")[-1].strip() if "|" in nxt else nxt
                    index += 1
            if not key:
                continue
            if value and not _is_document_stamp(value):
                out[key] = value
            elif key in access_keys:
                out.setdefault(key, "To confirm")
    return out


def extract_text_from_upload(filename: str, data: bytes) -> str:
    """Pull plain text from xlsx / csv / docx / pdf / txt."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _xlsx_to_text(data)
    if name.endswith(".csv"):
        return data.decode("utf-8", errors="ignore")
    if name.endswith(".docx"):
        return _docx_to_text(data)
    if name.endswith(".pdf"):
        return _pdf_to_text(data)
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")
    # try xlsx then utf-8
    try:
        return _xlsx_to_text(data)
    except Exception:  # noqa: BLE001
        return data.decode("utf-8", errors="ignore")


def _xlsx_to_text(data: bytes) -> str:
    cached = _xlsx_rows(data, data_only=True)
    # data_only is blank until Excel has cached formula results.
    if len(cached.strip()) < 40:
        return _xlsx_rows(data, data_only=False)
    return cached


def _xlsx_rows(data: bytes, *, data_only: bool) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=data_only)
    lines: list[str] = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines.append(f"## Sheet: {sheet}")
        for row in ws.iter_rows(values_only=True):
            vals = ["" if c is None else str(c).strip() for c in row]
            if any(vals):
                lines.append("|".join(vals))
    return "\n".join(lines)


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        t = page.extract_text() or ""
        if t.strip():
            parts.append(t)
    return "\n".join(parts)
