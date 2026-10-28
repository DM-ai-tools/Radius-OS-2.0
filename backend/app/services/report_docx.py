"""Render structured phase reports as a downloadable Word (.docx) document.

Reuses the same section blueprints, field normalization, and title resolution
as report_pdf.py / report_pdf_document.py — a blueprint fix (new section, new
column, a dropped-field fix) should apply to both formats without being
duplicated here. Only the actual document-building calls are format-specific.

The look matches the PDF: a deep-teal cover block, numbered section headings
with a teal rule, the summary in a mist panel, key figures as value-over-label
cells, light-header tables without vertical rules, amber only for warnings,
and a running header and "Page X of Y" footer after the cover.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.logging_config import get_logger
from app.services.report_export import resolve_report_title
from app.services.report_pdf_document import (
    SectionSpec,
    _get,
    _meta_line,
    _resolve_blueprint,
    _rows,
    _spec_key,
    _summary_text,
    _table_rows,
    auto_section_specs,
)
from app.services.report_pdf_normalize import normalize_report_payload

log = get_logger("report_docx")

INK = RGBColor(0x07, 0x1C, 0x1F)
MUTED = RGBColor(0x3D, 0x56, 0x5A)
PRIMARY = RGBColor(0x03, 0x56, 0x5F)
TEAL = RGBColor(0x02, 0x80, 0x90)
AMBER = RGBColor(0xB2, 0x6B, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT_TEAL = RGBColor(0x9F, 0xD8, 0xD2)

HEX_DEEP = "03565F"
HEX_TEAL = "028090"
HEX_MIST = "E6F5F3"
HEX_LINE = "D8E0E1"
HEX_AMBER_BG = "FAEEDA"
HEX_AMBER = "B26B00"
HEX_ZEBRA = "F7FAFA"

FONT = "Arial"
MAX_CELL_CHARS = 300
METRICS_PER_ROW = 4


# --------------------------------------------------------------------------- XML helpers

# Word rejects a file whose property elements are out of schema order, so every
# helper below re-sorts the parent after adding a child (ECMA-376 sequences).
_ORDER = {
    "pPr": [
        "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
        "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
        "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
        "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
        "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
        "sectPr", "pPrChange",
    ],
    "tcPr": [
        "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
        "textDirection", "tcFitText", "vAlign", "hideMark",
    ],
    "tblPr": [
        "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
        "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
    ],
}


def _reorder(parent: Any) -> None:
    order = _ORDER.get(parent.tag.split("}")[-1])
    if not order:
        return
    rank = {name: i for i, name in enumerate(order)}
    children = list(parent)
    children.sort(key=lambda el: rank.get(el.tag.split("}")[-1], len(order)))
    for el in children:
        parent.remove(el)
        parent.append(el)


def _shade(cell: Any, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(shd)
    _reorder(tc_pr)


def _cell_borders(cell: Any, **edges: tuple[int, str] | None) -> None:
    """edges: top/left/bottom/right=(eighths of a point, hex) or None for none."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1])
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tc_pr.append(borders)
    _reorder(tc_pr)


def _table_no_borders(table: Any) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)
    _reorder(tbl_pr)


def _cell_margins(cell: Any, top: int = 80, bottom: int = 80, left: int = 120, right: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)
    _reorder(tc_pr)


def _para_bottom_rule(paragraph: Any, *, size: int = 10, color: str = HEX_TEAL) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    p_pr.append(bdr)
    _reorder(p_pr)


def _keep_with_next(paragraph: Any) -> None:
    paragraph.paragraph_format.keep_with_next = True


def _field(paragraph: Any, instr: str) -> None:
    """A Word field (PAGE, NUMPAGES) that Word fills in when the file opens."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


# --------------------------------------------------------------------------- text helpers


def _cell_text(value: Any) -> str:
    if isinstance(value, tuple):
        return "\n".join(_cell_text(v) for v in value if v not in (None, ""))
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, int):
        text = f"{value:,}"
    elif isinstance(value, float):
        text = f"{int(value):,}" if value.is_integer() else f"{value:,.1f}"
    else:
        text = str(value)
    return text if len(text) <= MAX_CELL_CHARS else text[: MAX_CELL_CHARS - 1] + "…"


def _run(paragraph: Any, text: str, *, size: float = 10, bold: bool = False, color: RGBColor = INK) -> Any:
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    return run


def _add_muted_paragraph(doc: Document, text: str, *, size: int = 9, center: bool = False) -> None:
    p = doc.add_paragraph()
    _run(p, text, size=size, color=MUTED)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _setup(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.9))
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    for name in ("List Bullet", "List Number"):
        doc.styles[name].font.name = FONT
        doc.styles[name].font.size = Pt(9.5)


# --------------------------------------------------------------------------- blocks


def _section_heading(doc: Document, number: int | None, title: str, *, level: str = "section") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == "section" else 4)
    p.paragraph_format.space_after = Pt(10)
    if number is not None:
        _run(p, f"{number}   ", size=15, bold=True, color=TEAL)
    _run(p, title, size=21 if level == "report" else 15, bold=True)
    _para_bottom_rule(p, size=18 if level == "report" else 10)
    _keep_with_next(p)


def _panel(doc: Document, lines: list[tuple[str, bool, float, RGBColor]], *, fill: str, rule: str) -> None:
    """One shaded cell with a coloured left rule (summary, alert, group header)."""
    table = doc.add_table(rows=1, cols=1)
    _table_no_borders(table)
    cell = table.rows[0].cells[0]
    _shade(cell, fill)
    _cell_borders(cell, left=(24, rule))
    _cell_margins(cell, top=100, bottom=100, left=200, right=160)
    first = True
    for text, bold, size, color in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(3)
        _run(p, text, size=size, bold=bold, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _metrics(doc: Document, metrics: list[tuple[str, Any]]) -> None:
    per_row = max(1, min(METRICS_PER_ROW, len(metrics)))
    rows = (len(metrics) + per_row - 1) // per_row
    table = doc.add_table(rows=rows, cols=per_row)
    _table_no_borders(table)
    for i, (label, value) in enumerate(metrics):
        cell = table.rows[i // per_row].cells[i % per_row]
        _shade(cell, HEX_MIST)
        _cell_borders(cell, top=(16, HEX_TEAL))
        _cell_margins(cell, top=100, bottom=90, left=140, right=100)
        text = _cell_text(value)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        _run(p, text, size=13 if len(text) <= 14 else 10, bold=True)
        lp = cell.add_paragraph()
        lp.paragraph_format.space_after = Pt(0)
        _run(lp, str(label), size=8, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_table(doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
    """Light header, hairline rows, no vertical rules (matches the PDF)."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _table_no_borders(table)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        _shade(cell, HEX_MIST)
        _cell_borders(cell, bottom=(8, HEX_TEAL))
        _cell_margins(cell, top=60, bottom=60, left=90, right=90)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        _run(cell.paragraphs[0], h, size=8.5, bold=True, color=PRIMARY)
    header_row = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_row.append(repeat)
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row[: len(headers)]):
            cell = cells[i]
            if r_index % 2:
                _shade(cell, HEX_ZEBRA)
            _cell_borders(cell, bottom=(2, HEX_LINE))
            _cell_margins(cell, top=50, bottom=50, left=90, right=90)
            lines = _cell_text(val).split("\n")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _run(p, lines[0], size=8.5)
            for extra in lines[1:]:
                ep = cell.add_paragraph()
                ep.paragraph_format.space_after = Pt(0)
                _run(ep, extra, size=8, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# --------------------------------------------------------------------------- sections


def _render_docx_section(doc: Document, section_num: int, spec: SectionSpec, payload: dict[str, Any]) -> bool:
    has_content = False
    heading_added = False

    def _ensure_heading() -> None:
        nonlocal heading_added
        if not heading_added:
            _section_heading(doc, section_num, spec.title)
            heading_added = True

    if spec.alert_field and _get(payload, spec.alert_field):
        _ensure_heading()
        _panel(
            doc,
            [(spec.alert_label, True, 10, AMBER), (str(_get(payload, spec.alert_field)), False, 10, INK)],
            fill=HEX_AMBER_BG,
            rule=HEX_AMBER,
        )
        has_content = True

    intro = spec.intro
    if not intro and spec.intro_field:
        val = _get(payload, spec.intro_field)
        intro = str(val) if val is not None else None
    if intro:
        _ensure_heading()
        doc.add_paragraph(str(intro))
        has_content = True

    if spec.metrics:
        metrics = [(label, _get(payload, path)) for label, path in spec.metrics]
        metrics = [(label, val) for label, val in metrics if val is not None and val != ""]
        if metrics:
            _ensure_heading()
            _metrics(doc, metrics)
            has_content = True

    if spec.table_source:
        rows = _rows(_get(payload, spec.table_source))
        if rows and spec.prepare_rows:
            rows = spec.prepare_rows(rows)
        if rows:
            keys = spec.table_keys or []
            headers = spec.table_headers or [k.replace("_", " ").title() for k in keys]
            table_rows = _table_rows(rows, keys, spec.row_fn)
            if table_rows:
                _ensure_heading()
                _add_table(doc, headers, table_rows[: spec.max_rows])
                if len(table_rows) > spec.max_rows:
                    _add_muted_paragraph(doc, f"…and {len(table_rows) - spec.max_rows} more rows")
                has_content = True

    if spec.groups:
        g = spec.groups
        for group in _rows(_get(payload, g.source)):
            _ensure_heading()
            alert = group.get("tone") == "alert"
            lines = [(str(group.get("title") or ""), True, 10.5, INK)]
            if group.get("subtitle"):
                lines.append((str(group["subtitle"]), False, 9.5, INK))
            if group.get("detail"):
                lines.append((str(group["detail"]), False, 8.5, MUTED))
            _panel(doc, lines, fill=HEX_AMBER_BG if alert else HEX_MIST, rule=HEX_AMBER if alert else HEX_TEAL)
            rows = [r for r in group.get("rows") or [] if isinstance(r, list)]
            if rows:
                _add_table(doc, g.headers, rows[: g.max_rows])
            if group.get("more"):
                _add_muted_paragraph(doc, f"…and {group['more']} more keyword(s) in this cluster.")
            has_content = True

    for extra in spec.tables:
        rows = _rows(_get(payload, extra.source))
        table_rows = _table_rows(rows, extra.keys, extra.row_fn) if rows else []
        if table_rows:
            _ensure_heading()
            sub = doc.add_paragraph()
            _run(sub, extra.title, size=10.5, bold=True)
            _keep_with_next(sub)
            _add_table(doc, extra.headers, table_rows[: extra.max_rows])
            if len(table_rows) > extra.max_rows:
                _add_muted_paragraph(doc, f"…and {len(table_rows) - extra.max_rows} more rows")
            has_content = True

    if spec.bullets_source:
        raw = _get(payload, spec.bullets_source)
        items: list[str] = []
        if isinstance(raw, list):
            for x in raw:
                if isinstance(x, str):
                    items.append(x)
                elif isinstance(x, dict):
                    items.append(
                        "; ".join(f"{k}: {v}" for k, v in x.items() if not isinstance(v, (list, dict)))
                    )
        if items:
            _ensure_heading()
            for item in items:
                doc.add_paragraph(item, style="List Bullet")
            has_content = True

    if spec.numbered_source:
        raw = _get(payload, spec.numbered_source)
        if isinstance(raw, list) and raw:
            _ensure_heading()
            for item in raw:
                doc.add_paragraph(str(item), style="List Number")
            has_content = True

    if spec.note_fields:
        bits = [str(_get(payload, f)) for f in spec.note_fields if _get(payload, f)]
        if bits:
            _ensure_heading()
            note = doc.add_paragraph()
            _run(note, f"{spec.note_label or 'Note:'} ", size=10, bold=True)
            _run(note, " ".join(bits), size=10)
            has_content = True

    return has_content


def _render_docx_report(
    doc: Document,
    report: dict,
    client: dict,
    *,
    show_title: bool = True,
    contents: list[str] | None = None,
) -> None:
    raw_payload = report.get("payload") if isinstance(report.get("payload"), dict) else {}
    card_type = str(raw_payload.get("card_type") or report.get("card_type") or "")
    payload = normalize_report_payload(card_type, raw_payload)
    title = resolve_report_title(report, client)
    meta = _meta_line(payload, client)

    if show_title:
        _section_heading(doc, None, title, level="report")
        if meta:
            _add_muted_paragraph(doc, meta)
    elif "Geo:" in meta:
        # The cover and running header already name the client and site.
        _add_muted_paragraph(doc, f"Market: {meta.split('Geo:', 1)[1].strip()}")

    if payload.get("scorecard_line"):
        # One line, every audit group, before the narrative.
        glance = doc.add_paragraph()
        _run(glance, "At a glance: ", size=10, bold=True)
        _run(glance, str(payload["scorecard_line"]), size=10)
    summary = _summary_text(payload, client) or f"Structured phase report for {client.get('name') or 'the client'}."
    _panel(doc, [("Summary", True, 10, INK), (summary, False, 10.5, INK)], fill=HEX_MIST, rule=HEX_TEAL)

    blueprint = _resolve_blueprint(card_type, payload)
    section_num = 0
    for spec in blueprint:
        if _render_docx_section(doc, section_num + 1, spec, payload):
            section_num += 1
            if contents is not None:
                contents.append(f"{section_num}. {spec.title}")

    # The same automatic sections the PDF adds (labelled, suppression-aware), so the
    # two exports agree. A report the PDF draws with its own renderer has none.
    from app.services.report_pdf import SPECIAL_RENDERERS

    if card_type in SPECIAL_RENDERERS:
        return
    rendered: list[str] = []
    for spec in auto_section_specs(card_type, payload, blueprint):
        if _render_docx_section(doc, section_num + 1, spec, payload):
            section_num += 1
            rendered.append(_spec_key(spec))
            if contents is not None:
                contents.append(f"{section_num}. {spec.title}")
    if rendered:
        log.warning("report_auto_sections_rendered", card_type=card_type or "(none)", keys=rendered, format="docx")


# --------------------------------------------------------------------------- cover + frame


def _prepared(bundle: dict) -> str:
    exported = bundle.get("exported_at") or ""
    if not exported:
        return datetime.now(UTC).strftime("%d %B %Y")
    try:
        return datetime.fromisoformat(str(exported).replace("Z", "+00:00")).strftime("%d %B %Y")
    except ValueError:
        return str(exported)


def _cover(
    doc: Document,
    *,
    title: str,
    subtitle: str,
    client_name: str,
    client_url: str,
    prepared: str,
    contents: list[str],
    contents_label: str,
) -> None:
    block = doc.add_table(rows=1, cols=1)
    _table_no_borders(block)
    cell = block.rows[0].cells[0]
    _shade(cell, HEX_DEEP)
    _cell_borders(cell, bottom=(48, HEX_TEAL))
    _cell_margins(cell, top=500, bottom=500, left=500, right=400)
    p = cell.paragraphs[0]
    _run(p, "Radius OS", size=10, bold=True, color=SOFT_TEAL)
    p.paragraph_format.space_after = Pt(150)
    tp = cell.add_paragraph()
    tp.paragraph_format.space_after = Pt(6)
    _run(tp, title, size=30, bold=True, color=WHITE)
    for text, size, bold in ((client_name, 13, True), (subtitle, 11, False), (client_url, 11, False)):
        if text:
            cp = cell.add_paragraph()
            cp.paragraph_format.space_after = Pt(2)
            _run(cp, text, size=size, bold=bold, color=WHITE if bold else SOFT_TEAL)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    details = [(label, value) for label, value in (("Prepared for", client_name), ("Website", client_url), ("Date", prepared)) if value]
    if details:
        table = doc.add_table(rows=2, cols=len(details))
        _table_no_borders(table)
        for i, (label, value) in enumerate(details):
            lp = table.rows[0].cells[i].paragraphs[0]
            lp.paragraph_format.space_after = Pt(0)
            _run(lp, label, size=8.5, color=MUTED)
            vp = table.rows[1].cells[i].paragraphs[0]
            _run(vp, value, size=10.5)
    if contents:
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(18)
        _run(head, contents_label, size=10, bold=True, color=PRIMARY)
        _para_bottom_rule(head, size=4, color=HEX_LINE)
        for entry in contents:
            item = doc.add_paragraph()
            item.paragraph_format.space_after = Pt(2)
            _run(item, entry, size=10)
    foot = doc.add_paragraph()
    foot.paragraph_format.space_before = Pt(24)
    _run(foot, "Prepared with Radius OS. Figures come from the data providers named in each section.", size=8, color=MUTED)


def _frame(doc: Document, *, left: str, right: str, footer: str) -> None:
    """Body section after the cover: running header, footer with Page X of Y."""
    section = doc.sections[-1]
    section.different_first_page_header_footer = False
    usable = section.page_width - section.left_margin - section.right_margin
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    _run(hp, left, size=8, bold=True, color=PRIMARY)
    _run(hp, f"\t{right}", size=8, color=MUTED)
    _para_bottom_rule(hp, size=4, color=HEX_LINE)
    footer_part = section.footer
    footer_part.is_linked_to_previous = False
    fp = footer_part.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    _run(fp, f"{footer}\tPage ", size=8, color=MUTED)
    _field(fp, "PAGE")
    _run(fp, " of ", size=8, color=MUTED)
    _field(fp, "NUMPAGES")


def _start_body(doc: Document) -> None:
    """New section for the body so the cover has no header or footer."""
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False


def build_reports_docx(bundle: dict, *, single: bool = False) -> bytes:
    """Build a Word document containing all client reports, or one when single=True."""
    client = bundle.get("client") or {}
    reports: list[dict] = bundle.get("reports") or []
    client_name = str(client.get("name") or "Client")
    single_title = resolve_report_title(reports[0], client) if single and reports else ""
    doc_title = single_title or f"{client_name}: phase reports"

    # The cover lists the sections, which are only known once rendered.
    contents: list[str] = []
    if single and reports:
        scratch = Document()
        _setup(scratch)
        _render_docx_report(scratch, reports[0], client, show_title=False, contents=contents)
        contents_label = "In this report"
    else:
        contents = [f"{i}. {resolve_report_title(r, client)}" for i, r in enumerate(reports, 1)]
        contents_label = f"{len(reports)} phase reports in this pack"
    readiness = (bundle.get("phase_statuses") or {}).get("overall_readiness_score")
    subtitle = str(client.get("industry") or "")
    if not single and readiness is not None:
        subtitle = f"Readiness {readiness:.0f}%" + (f", {subtitle}" if subtitle else "")

    doc = Document()
    _setup(doc)
    doc.core_properties.title = doc_title
    doc.core_properties.author = "Radius OS"
    _cover(
        doc,
        title=single_title or "Phase reports",
        subtitle=subtitle,
        client_name=client_name,
        client_url=str(client.get("primary_url") or ""),
        prepared=_prepared(bundle),
        contents=contents,
        contents_label=contents_label,
    )
    _start_body(doc)
    _frame(
        doc,
        left=doc_title,
        right=str(client.get("primary_url") or client_name),
        footer=f"Prepared with Radius OS for {client_name}",
    )

    if single and reports:
        _render_docx_report(doc, reports[0], client, show_title=False)
    else:
        for i, report in enumerate(reports):
            if i > 0:
                doc.add_page_break()
            _render_docx_report(doc, report, client, show_title=True)

    if not reports:
        doc.add_paragraph("No reports generated yet.")
        doc.add_paragraph(
            "Run discovery, tracking, website audit, and downstream phases in chat, then download again."
        )

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
