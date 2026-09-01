"""BW workbook row builders — match ``BW_Category_Mapping_Search_Demand_Analysis.xlsx``.

Sheets:
- Category & URL Mapping (Phase 6 / site architecture)
- Search Demand Analysis (Phase 5 / search demand)
- TOFU & MOFU Content Strategy (content strategy)
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from app.services.url_mapping import resolve_sheet_url_columns

CATEGORY_URL_COLUMNS = [
    "level",
    "l1_category",
    "l2_subcategory",
    "l3_sub_subcategory",
    "l4_attribution",
    "current_url",
    "proposed_url",
    "status",
    "primary_keyword",
    "search_volume_mo",
    "cpc",
    "secondary_keywords",
    "combined_cluster_volume",
    "page_type",
    "priority",
    "notes",
    "est_products",
    "in_scope",
]

SEARCH_DEMAND_COLUMNS = [
    "keyword",
    "search_volume_mo",
    "cpc",
    "competition",
    "category_dimension",
    "parent_category",
    "target_url",
    "status",
    "funnel_stage",
    "secondary_keywords",
    "combined_cluster_volume",
    "priority",
]

TOFU_MOFU_COLUMNS = [
    "content_type",
    "funnel_stage",
    "parent_category",
    "blog_title",
    "primary_keyword",
    "search_volume_mo",
    "cpc",
    "secondary_keywords",
    "combined_cluster_volume",
    "internal_linking_targets",
    "priority",
    "notes",
]


def _fmt_level(level: Any) -> str:
    try:
        n = int(level or 0)
    except (TypeError, ValueError):
        return ""
    return f"L{n}" if n else ""


def _excel_status(action: str | None, status: str | None = None) -> str:
    raw = str(status or action or "").upper()
    if raw in ("LIVE", "NEW"):
        return raw
    if "CREATE" in raw or raw == "CREATE NEW":
        return "NEW"
    if raw in ("OPTIMIZE_EXISTING", "REVIEW_MERGE_REDIRECT", "OPTIMIZE EXISTING", "REVIEW / MERGE / REDIRECT"):
        return "LIVE"
    return "NEW" if "create" in raw.lower() else "LIVE"


def _excel_priority(
    *,
    band: str | None = None,
    priority: str | None = None,
    volume: Any = None,
    page_type: str | None = None,
) -> str:
    if priority and str(priority).upper().startswith("P"):
        return str(priority).upper()
    band_u = str(band or priority or "").upper()
    if band_u in ("P0", "P1", "P2", "P3"):
        return band_u
    if band_u in ("HIGH", "QUICK WIN", "BIG BET"):
        return "P0" if page_type and "hub" in str(page_type).lower() else "P1"
    if band_u in ("MEDIUM", "FILL-IN"):
        return "P2"
    if band_u in ("LOW", "AVOID"):
        return "P3"
    try:
        vol = int(volume or 0)
    except (TypeError, ValueError):
        vol = 0
    if vol >= 20000:
        return "P0"
    if vol >= 5000:
        return "P1"
    if vol >= 1000:
        return "P2"
    return "P3"


def _join_keywords(raw: Any, limit: int = 10) -> str:
    if isinstance(raw, str):
        return raw
    if not isinstance(raw, list):
        return ""
    return ", ".join(str(x).strip() for x in raw[:limit] if str(x).strip())


def _num(value: Any) -> int | float | None:
    if value is None or value == "":
        return None
    try:
        if isinstance(value, float):
            return value
        if isinstance(value, int):
            return value
        text = str(value).replace(",", "").strip()
        if "." in text:
            return float(text)
        return int(text)
    except (TypeError, ValueError):
        return None


def build_category_url_mapping_rows(
    url_map_entries: list[dict[str, Any]],
    *,
    client_name: str = "",
    include_section_headers: bool = True,
) -> list[dict[str, Any]]:
    """Phase 6 rows — Category & URL Mapping sheet."""
    rows: list[dict[str, Any]] = []
    last_l1: str | None = None
    sorted_entries = sorted(
        url_map_entries,
        key=lambda e: (
            str(e.get("l1_category") or e.get("pillar") or ""),
            str(e.get("l2_subcategory") or ""),
            str(e.get("l3_subsubcategory") or e.get("l3_sub_subcategory") or ""),
            str(e.get("primary_keyword") or ""),
        ),
    )
    for entry in sorted_entries:
        if not isinstance(entry, dict):
            continue
        l1 = str(entry.get("l1_category") or entry.get("pillar") or "").strip()
        if include_section_headers and l1 and l1 != last_l1:
            rows.append(
                {
                    "row_type": "section",
                    "section_label": f"▶ {l1.upper()}",
                    "l1_category": l1,
                }
            )
            last_l1 = l1
        action = str(entry.get("action") or "")
        status = _excel_status(action, entry.get("status"))
        page_type = str(entry.get("page_type") or "")
        volume = entry.get("search_volume") or entry.get("volume")
        current_url, proposed_url = resolve_sheet_url_columns(
            current_url=entry.get("current_url"),
            proposed_url=entry.get("proposed_url"),
            action=action,
            selected_url=entry.get("selected_url"),
            create_url=entry.get("create_url"),
        )
        rows.append(
            {
                "row_type": "data",
                "level": _fmt_level(entry.get("level")),
                "l1_category": l1 or None,
                "l2_subcategory": entry.get("l2_subcategory"),
                "l3_sub_subcategory": entry.get("l3_subsubcategory")
                or entry.get("l3_sub_subcategory"),
                "l4_attribution": entry.get("l4_attribution"),
                "current_url": current_url,
                "proposed_url": proposed_url,
                "status": status,
                "primary_keyword": entry.get("primary_keyword"),
                "search_volume_mo": volume,
                "cpc": entry.get("cpc"),
                "secondary_keywords": _join_keywords(
                    entry.get("secondary_keywords_sheet") or entry.get("secondary_keywords")
                ),
                "combined_cluster_volume": entry.get("combined_cluster_volume"),
                "page_type": page_type,
                "priority": _excel_priority(
                    band=entry.get("score_band") or entry.get("priority"),
                    priority=entry.get("priority"),
                    volume=volume,
                    page_type=page_type,
                ),
                "notes": entry.get("notes"),
                "est_products": entry.get("est_products"),
                "in_scope": entry.get("in_scope") if entry.get("in_scope") is not None else "YES",
                "cluster": entry.get("cluster"),
                "url_score": entry.get("url_score"),
            }
        )
    return rows


def _target_url_by_keyword(url_map: list[dict[str, Any]]) -> dict[str, str]:
    out: dict[str, str] = {}
    for row in url_map:
        if not isinstance(row, dict):
            continue
        kw = str(row.get("primary_keyword") or "").strip().lower()
        if not kw:
            continue
        url = row.get("current_url") or row.get("proposed_url") or row.get("selected_url")
        if url:
            out[kw] = str(url)
    return out


def _url_status_by_keyword(url_map: list[dict[str, Any]]) -> dict[str, str]:
    out: dict[str, str] = {}
    for row in url_map:
        if not isinstance(row, dict):
            continue
        kw = str(row.get("primary_keyword") or "").strip().lower()
        if not kw:
            continue
        out[kw] = _excel_status(str(row.get("action") or ""), row.get("status"))
    return out


def build_search_demand_analysis_rows(
    keyword_dataset: list[dict[str, Any]],
    *,
    clusters: list[dict[str, Any]] | None = None,
    url_map: list[dict[str, Any]] | None = None,
    limit: int = 500,
) -> list[dict[str, Any]]:
    """Phase 5 rows — Search Demand Analysis sheet."""
    cluster_by_kw: dict[str, dict[str, Any]] = {}
    for cluster in clusters or []:
        if not isinstance(cluster, dict):
            continue
        pkw = str(cluster.get("primary_keyword") or "").strip().lower()
        if pkw:
            cluster_by_kw[pkw] = cluster
        for row in cluster.get("keywords") or []:
            if isinstance(row, dict):
                kw = str(row.get("keyword") or "").strip().lower()
                if kw:
                    cluster_by_kw.setdefault(kw, cluster)

    targets = _target_url_by_keyword(url_map or [])
    url_statuses = _url_status_by_keyword(url_map or [])
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    dataset = sorted(
        [r for r in keyword_dataset if isinstance(r, dict) and r.get("keyword")],
        key=lambda r: -int(_num(r.get("volume")) or 0),
    )
    for row in dataset[:limit]:
        kw = str(row.get("keyword") or "").strip()
        key = kw.lower()
        if not kw or key in seen:
            continue
        seen.add(key)
        cluster = cluster_by_kw.get(key) or {}
        parent = (
            row.get("parent_topic")
            or row.get("target")
            or cluster.get("name")
            or row.get("seed")
        )
        category_dim = row.get("target_type") or cluster.get("name") or parent
        target_url = targets.get(key) or row.get("recommended_url") or cluster.get("recommended_url")
        funnel = str(row.get("funnel") or "BOFU").upper()
        secondaries = row.get("secondary_keywords") or row.get("supporting_keywords") or []
        if not secondaries and cluster:
            from app.services.content_pipeline import identify_secondary_keywords

            secondaries = identify_secondary_keywords(cluster, limit=10)
        combined = row.get("combined_cluster_volume")
        if combined is None and cluster:
            from app.services.content_pipeline import combined_cluster_volume

            combined = combined_cluster_volume(cluster)
        difficulty = row.get("difficulty")
        competition = difficulty
        if isinstance(difficulty, (int, float)) and difficulty <= 1:
            competition = round(float(difficulty) * 100)
        status = url_statuses.get(key) or ("LIVE" if target_url and key in targets else "NEW")
        rows.append(
            {
                "row_type": "data",
                "keyword": kw,
                "search_volume_mo": row.get("volume"),
                "cpc": row.get("cpc"),
                "competition": competition,
                "category_dimension": category_dim,
                "parent_category": parent,
                "target_url": target_url or "",
                "status": status,
                "funnel_stage": funnel,
                "secondary_keywords": _join_keywords(secondaries),
                "combined_cluster_volume": combined,
                "priority": _excel_priority(volume=row.get("volume"), band=row.get("bucket")),
                "intent": row.get("intent"),
                "opportunity_score": row.get("opportunity_score"),
            }
        )
    return rows


def build_tofu_mofu_content_strategy_rows(
    priority_queue: list[dict[str, Any]],
    *,
    site_architecture: dict[str, Any] | None = None,
    limit: int = 100,
) -> list[dict[str, Any]]:
    """Content strategy rows — TOFU & MOFU sheet only (excludes BOFU product pages)."""
    url_map = list(
        ((site_architecture or {}).get("url_map_report") or {}).get("final_url_map") or []
    )
    hub_urls = [
        str(r.get("current_url") or r.get("selected_url") or "")
        for r in url_map
        if str(r.get("page_type") or "").lower().find("hub") >= 0
    ][:5]
    rows: list[dict[str, Any]] = []
    for item in priority_queue:
        if not isinstance(item, dict):
            continue
        funnel = str(item.get("funnel") or "").upper()
        if funnel not in ("TOFU", "MOFU"):
            continue
        if len(rows) >= limit:
            break
        kw = str(item.get("primary_keyword") or item.get("keyword") or "").strip()
        secondaries = item.get("secondary_keywords") or item.get("supporting_keywords") or []
        internal = item.get("internal_linking_targets")
        if not internal:
            internal = hub_urls or [
                str(item.get("suggested_url") or ""),
            ]
        rows.append(
            {
                "row_type": "data",
                "content_type": item.get("content_type") or "Blog",
                "funnel_stage": funnel,
                "parent_category": item.get("parent_category")
                or item.get("pillar")
                or item.get("cluster")
                or item.get("parent"),
                "blog_title": item.get("title"),
                "primary_keyword": kw,
                "search_volume_mo": item.get("volume"),
                "cpc": item.get("cpc"),
                "secondary_keywords": _join_keywords(secondaries),
                "combined_cluster_volume": item.get("combined_cluster_volume"),
                "internal_linking_targets": _join_keywords(
                    internal if isinstance(internal, list) else str(internal).split(",")
                ),
                "priority": _excel_priority(
                    priority=item.get("priority"),
                    volume=item.get("volume"),
                ),
                "notes": item.get("rationale") or item.get("priority_note") or item.get("angle"),
                "suggested_url": item.get("suggested_url"),
            }
        )
    return rows


def attach_workbook_to_search_demand(summary: dict[str, Any]) -> dict[str, Any]:
    summary = dict(summary)
    summary["search_demand_analysis"] = build_search_demand_analysis_rows(
        list(summary.get("keyword_dataset") or []),
        clusters=list(summary.get("clusters") or []),
        url_map=[],
    )
    summary["workbook"] = {
        "sheet": "Search Demand Analysis",
        "columns": SEARCH_DEMAND_COLUMNS,
        "rows": summary["search_demand_analysis"],
        "row_count": len(summary["search_demand_analysis"]),
    }
    return summary


def attach_workbook_to_site_architecture(summary: dict[str, Any], *, client_name: str = "") -> dict[str, Any]:
    summary = dict(summary)
    url_map = list(
        ((summary.get("url_map_report") or {}).get("final_url_map") or [])
    )
    summary["category_url_mapping"] = build_category_url_mapping_rows(
        url_map,
        client_name=client_name,
    )
    summary["workbook"] = {
        "sheet": "Category & URL Mapping",
        "columns": CATEGORY_URL_COLUMNS,
        "rows": summary["category_url_mapping"],
        "row_count": len(summary["category_url_mapping"]),
    }
    return summary


def attach_workbook_to_content_strategy(
    summary: dict[str, Any],
    *,
    site_architecture: dict[str, Any] | None = None,
) -> dict[str, Any]:
    summary = dict(summary)
    summary["tofu_mofu_content_strategy"] = build_tofu_mofu_content_strategy_rows(
        list(summary.get("priority_queue") or []),
        site_architecture=site_architecture,
    )
    summary["workbook"] = {
        "sheet": "TOFU & MOFU Content Strategy",
        "columns": TOFU_MOFU_COLUMNS,
        "rows": summary["tofu_mofu_content_strategy"],
        "row_count": len(summary["tofu_mofu_content_strategy"]),
    }
    return summary


def build_combined_workbook_pack(
    *,
    client_name: str,
    search_demand: dict[str, Any] | None = None,
    site_architecture: dict[str, Any] | None = None,
    content_strategy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    sd = dict(search_demand or {})
    ia = dict(site_architecture or {})
    cs = dict(content_strategy or {})
    url_map = list(((ia.get("url_map_report") or {}).get("final_url_map") or []))
    return {
        "client_name": client_name,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "template": "BW_Category_Mapping_Search_Demand_Analysis",
        "sheets": {
            "category_url_mapping": {
                "title": "Category & URL Mapping",
                "columns": CATEGORY_URL_COLUMNS,
                "rows": build_category_url_mapping_rows(url_map, client_name=client_name),
            },
            "search_demand_analysis": {
                "title": "Search Demand Analysis",
                "columns": SEARCH_DEMAND_COLUMNS,
                "rows": build_search_demand_analysis_rows(
                    list(sd.get("keyword_dataset") or []),
                    clusters=list(sd.get("clusters") or []),
                    url_map=url_map,
                ),
            },
            "tofu_mofu_content_strategy": {
                "title": "TOFU & MOFU Content Strategy",
                "columns": TOFU_MOFU_COLUMNS,
                "rows": build_tofu_mofu_content_strategy_rows(
                    list(cs.get("priority_queue") or []),
                    site_architecture=ia,
                ),
            },
        },
    }


def export_workbook_xlsx(pack: dict[str, Any]) -> bytes:
    """Export combined workbook pack to .xlsx bytes."""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True)
    for sheet_key, sheet in (pack.get("sheets") or {}).items():
        title = str(sheet.get("title") or sheet_key)
        ws = wb.create_sheet(title[:31])
        columns = list(sheet.get("columns") or [])
        ws.append([c.replace("_", " ").title() for c in columns])
        for cell in ws[1]:
            cell.font = header_font
        for row in sheet.get("rows") or []:
            if not isinstance(row, dict):
                continue
            if row.get("row_type") == "section":
                ws.append([row.get("section_label")] + [""] * (len(columns) - 1))
                continue
            ws.append([row.get(col) for col in columns])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
