"""Cluster structure and the layered Keyword Clusters report (2026-10-19 brief).

- Oversized clusters split by modifier pattern, then intent, then funnel, and
  the split is named; merges and splits are counted.
- Every cluster carries its page pattern; comparison and location stand out.
- Pillar hubs group clusters sharing a head topic that together span the funnel.
- New fields survive clusters_for_cdp.
- The report shows summary, hubs, overview, detail, orphans, roadmap, quality
  and provenance, and says what was verified versus inferred.
"""

from __future__ import annotations

from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest

from app.services.keyword_cluster_report import DETAIL_CLUSTER_CAP, build_cluster_report_view
from app.services.keyword_cluster_roles import RoleContext
from app.services.keyword_cluster_structure import (
    SPLIT_THRESHOLD,
    assign_pillars,
    cluster_modifier_pattern,
    pattern_of,
    split_name,
    split_oversized,
)
from app.services.keyword_clustering import (
    build_clusters_deterministic,
    clusters_for_cdp,
    run_keyword_clustering,
)
from app.services.keyword_llm_clustering import build_cluster_report_from_llm


def _row(kw, vol=100, intent="informational", parent=None, **extra):
    return {
        "keyword": kw, "volume": vol, "difficulty": 20, "intent": intent, "provider_intent": intent,
        "intent_basis": "measured", "parent_topic": parent, "relevance_reason": "seed_phrase", **extra,
    }


OVERSIZED = (
    [_row(f"local seo for {a}", parent="local seo") for a in
     ("dentists", "lawyers", "plumbers", "builders", "accountants", "startups", "cafes", "gyms",
      "doctors", "real estate", "electricians", "salons")]
    + [_row(f"local seo {t}", parent="local seo") for t in
       ("tools", "software", "platform", "app", "plugin tools", "audit software")]
    + [_row(f"local seo vs {x}", parent="local seo") for x in ("ppc", "sem", "social", "ads")]
    + [_row("what is local seo", parent="local seo"), _row("local seo guide", parent="local seo")]
)


# --------------------------------------------------------------- size discipline


def test_split_goes_by_modifier_first_and_loses_nothing():
    assert len(OVERSIZED) > SPLIT_THRESHOLD
    parts = split_oversized(OVERSIZED, lambda r: (r["intent"], "TOFU"))
    assert len(parts) >= 2
    assert sorted(k["keyword"] for part, _p in parts for k in part) == sorted(r["keyword"] for r in OVERSIZED)
    dims = {path[0][0] for _part, path in parts}
    assert dims == {"modifier"}
    assert all(len(part) >= 3 for part, _p in parts)
    names = {split_name("Local SEO", path) for _part, path in parts}
    assert "Local SEO: audience-specific" in names and "Local SEO: comparisons" in names


def test_split_falls_back_to_intent_then_funnel():
    same_pattern = [_row(f"seo topic {i}", intent="informational" if i % 2 else "commercial") for i in range(18)]
    parts = split_oversized(same_pattern, lambda r: (r["intent"], "MOFU"))
    assert {path[0][0] for _p, path in parts} == {"intent"}
    one_intent = [_row(f"seo topic {i}") for i in range(18)]
    parts = split_oversized(one_intent, lambda r: ("informational", "TOFU" if r["keyword"][-1] in "02468" else "MOFU"))
    assert {path[0][0] for _p, path in parts} == {"funnel"}
    unsplittable = split_oversized(one_intent, lambda r: ("informational", "TOFU"))
    assert unsplittable == [(one_intent, [])]


def test_rule_path_keeps_a_primary_with_its_secondaries_and_balances():
    report = build_clusters_deterministic([dict(r) for r in OVERSIZED])
    by_kw = {
        row["keyword"]: cluster["primary_keyword"]
        for cluster in report["clusters"]
        for row in cluster["keywords"]
    }
    assert by_kw["local seo for dentists"] != by_kw["local seo for lawyers"]
    assert report["keyword_accounting"]["balanced"] is True
    assert report["keyword_accounting"]["input_count"] == len(OVERSIZED)
    for cluster in report["clusters"]:
        roles = {row["role"] for row in cluster["keywords"]}
        assert roles <= {"Primary", "Secondary"}
        assert cluster["keyword_count"] <= SPLIT_THRESHOLD


def test_model_path_applies_the_same_split():
    rows = [dict(r) for r in OVERSIZED]
    payload = {"clusters": [{"name": "Local SEO", "keywords": [{"keyword": r["keyword"]} for r in rows]}]}
    report = build_cluster_report_from_llm(payload, rows, role_context=RoleContext.build(rows))
    parts = [c for c in report["clusters"] if c.get("split_from")]
    assert len(parts) >= 2
    assert report["size_discipline"]["split_clusters"] == 1
    assert sum(c["keyword_count"] for c in report["clusters"]) == len(rows)


# --------------------------------------------------------------- page pattern


def test_page_pattern_is_read_from_members_and_locations_are_whole_words():
    assert pattern_of("seo vs ppc") == ("comparison", None)
    assert pattern_of("seo agency sydney") == ("location", "sydney")
    assert pattern_of("seo cost australia") == ("location", "australia")  # not "us"
    assert pattern_of("business seo") == ("core", None)  # "us" inside "business" is no place
    assert pattern_of("seo services") == ("core", None)
    cluster = {"primary_keyword": "seo vs ppc", "keywords": [{"keyword": "seo vs ppc"}, {"keyword": "seo vs sem"}]}
    assert cluster_modifier_pattern(cluster)["modifier_label"] == "Comparison"
    near = {"primary_keyword": "seo near me", "keywords": [{"keyword": "seo near me"}]}
    assert cluster_modifier_pattern(near)["modifier_label"] == "Local (near me)"


# --------------------------------------------------------------- pillars


def _cluster(name, primary, funnel, volume=100, service=None):
    return {"name": name, "primary_keyword": primary, "funnel": funnel, "intent": "commercial",
            "est_traffic": volume, "service": service}


def test_pillars_need_a_shared_head_two_subtopics_and_more_than_one_stage():
    clusters = [
        _cluster("What Is SEO", "what is seo", "TOFU", 9000),
        _cluster("SEO Cost", "how much does seo cost", "MOFU", 1000),
        _cluster("SEO Agency", "seo agency melbourne", "BOFU", 1300),
        _cluster("Google Ads Agency", "google ads agency", "BOFU"),
        _cluster("Google Ads Management", "google ads management", "BOFU"),
        _cluster("Podcast", "marketing podcast", "TOFU"),
    ]
    pillars = assign_pillars(clusters)
    assert [p["name"] for p in pillars] == ["SEO"]
    seo = pillars[0]
    assert seo["subtopic_count"] == 3 and seo["total_volume"] == 11300
    assert seo["funnel_stages"] == ["TOFU", "MOFU", "BOFU"]
    by_name = {c["name"]: c for c in clusters}
    assert by_name["SEO Agency"]["cluster_tier"] == "subtopic" and by_name["SEO Agency"]["pillar_name"] == "SEO"
    # Same head, but one stage only: not a hub, and the reason says so.
    assert by_name["Google Ads Agency"]["cluster_tier"] == "standalone"
    assert "BOFU" in by_name["Google Ads Agency"]["standalone_reason"]
    assert by_name["Podcast"]["cluster_tier"] == "standalone"


def test_pillars_do_not_cross_services():
    clusters = [
        _cluster("A", "seo tips", "TOFU", service="SEO"),
        _cluster("B", "seo agency", "BOFU", service="Local SEO"),
    ]
    assert assign_pillars(clusters, service_of=lambda c: c["service"]) == []


# --------------------------------------------------------------- full run + CDP


ROWS = [
    _row("what is seo", 9900, "informational", service="SEO"),
    _row("seo tips", 720, "informational", service="SEO"),
    _row("how much does seo cost", 1000, "commercial", service="SEO"),
    _row("seo vs ppc", 880, "commercial", service="SEO"),
    _row("seo agency melbourne", 1300, "commercial", service="SEO"),
    _row("seo company melbourne", 590, "commercial", service="SEO"),
    _row("marketing quotes", None, None, service="SEO", relevance_reason="theme_overlap"),
]


@pytest.mark.asyncio
async def test_full_rule_run_carries_structure_and_cdp_keeps_it():
    report = await run_keyword_clustering([dict(r) for r in ROWS], within_deadline=False)
    assert report["keyword_accounting"]["balanced"] is True
    assert {"pillars", "size_discipline", "clustering_provenance", "content_roadmap"} <= set(report)
    prov = report["clustering_provenance"]
    assert prov["grouping"] == "rules" and prov["serp_check_ran"] is False
    assert any("inferred" in t or "did not run" in t for t in prov["inferred"])
    for cluster in report["clusters"]:
        assert cluster["modifier_label"] and cluster["cluster_tier"] in ("subtopic", "standalone")
    flat = clusters_for_cdp(report)
    for key in ("modifier_pattern", "modifier_label", "cluster_tier", "pillar_name", "split_from",
                "primary_selection", "intent_distribution", "role_counts"):
        assert key in flat[0]
    assert [r["priority"] for r in report["content_roadmap"]] == list(range(1, len(report["content_roadmap"]) + 1))
    assert {"intent", "funnel", "target_keyword", "content_type"} <= set(report["content_roadmap"][0])


def test_model_dropped_keywords_become_orphans_with_their_reason():
    rows = [dict(r) for r in ROWS]
    payload = {
        "dropped": [{"keyword": "marketing quotes", "reason": "not a service search"}],
        "clusters": [{"name": "SEO", "keywords": [{"keyword": r["keyword"]} for r in rows[:6]]}],
    }
    report = build_cluster_report_from_llm(payload, rows, role_context=RoleContext.build(rows))
    orphan = next(o for o in report["orphans"] if o["keyword"] == "marketing quotes")
    assert orphan["exclusion_reason"] == "dropped_for_relevance"
    clustered = sum(c["keyword_count"] for c in report["clusters"])
    assert clustered + len(report["orphans"]) == len(rows)


def test_serp_merge_carries_the_folded_cluster_size():
    from app.services.keyword_serp_clustering import merge_clusters_by_serp_overlap

    a = {"primary_keyword": "seo agency", "keyword_count": 30, "keywords": [{"keyword": "seo agency"}]}
    b = {"primary_keyword": "seo company", "keyword_count": 12, "keywords": [{"keyword": "seo company"}]}
    merged = merge_clusters_by_serp_overlap([a, b], [["seo agency", "seo company"]])
    assert merged[0]["keyword_count"] == 42
    assert merged[0]["serp_group"] == ["seo agency", "seo company"]


# --------------------------------------------------------------- report view


def _report(**overrides):
    clusters = []
    for i in range(DETAIL_CLUSTER_CAP + 3):
        clusters.append({
            "name": f"Cluster {i}", "primary_keyword": f"kw {i}", "intent": "commercial", "funnel": "MOFU",
            "content_type": "guide", "recommended_content": "Consideration Guide", "recommended_url": f"/blog/kw-{i}",
            "keyword_count": 2, "est_traffic": 100 + i, "avg_difficulty": 20.0,
            "modifier_pattern": "comparison" if i == 0 else "core",
            "modifier_label": "Comparison" if i == 0 else "Core topic",
            "primary_selection": {"reason": "Exact match for the service."},
            "keywords": [
                {"keyword": f"kw {i}", "role": "Primary", "volume": 100, "difficulty": 20, "intent": "commercial", "funnel": "MOFU", "opportunity_score": 50},
                {"keyword": f"kw {i} b", "role": "Supporting", "volume": None, "difficulty": None, "intent": "informational", "funnel": "TOFU", "opportunity_score": None},
            ],
        })
    report = {
        "clusters": clusters,
        "total_keywords": 40,
        "keyword_accounting": {"input_count": 40, "clustered_count": 30, "orphan_count": 8, "unaccounted_count": 2, "balanced": False},
        "orphans": [
            {"keyword": "seo podcast", "volume": None, "intent": "informational", "exclusion_reason": "no_topic_fit"},
            {"keyword": "marketing quotes", "volume": 50, "intent": "informational", "exclusion_reason": "dropped_for_relevance", "notes": "not a service search"},
        ],
        "orphan_count": 8,
        "content_roadmap": [{"priority": 2, "cluster": "B"}, {"priority": 1, "cluster": "A"}],
        "cluster_validation": {"issues": [
            {"cluster": "Cluster 3", "issue": "single_keyword_cluster", "severity": "warn"},
            {"cluster": "Cluster 4", "issue": "single_keyword_cluster", "severity": "warn"},
            {"cluster": "Cluster 2", "issue": "duplicate_primary_keyword", "keyword": "kw 2"},
            {"cluster": "Cluster 5", "issue": "missing_content_type"},
            {"cluster": "Cluster 1", "issue": "mixed_intent_cluster", "intent_distribution": {"informational": 40.0, "transactional": 35.0}},
        ]},
        "size_discipline": {"split_threshold": 15, "merged_small_clusters": 2, "split_clusters": 1, "clusters_from_splits": 3},
        "pillars": [{"name": "SEO", "recommended_url": "/seo/", "subtopic_count": 2, "total_volume": 1500,
                     "funnel_stages": ["TOFU", "BOFU"], "subtopics": [{"name": "Cluster 0", "volume": 100}]}],
        "clustering_provenance": {
            "grouping": "model", "source_label": "Model clustering, SERP check skipped", "serp_check_ran": False,
            "serp_check_reason": "locale_unresolved", "naming_ran": False, "model_clustering_ran": True,
            "model_clustering_model": "claude-sonnet-5", "invented_keywords_rejected": 3,
            "cleaning_excluded_count": 12, "cleaning_excluded_by_reason": {"stale_year": 5, "competitor_brand": 7},
            "verified": ["placed keywords checked"], "inferred": ["no SERP check"],
        },
    }
    report.update(overrides)
    return report


def test_view_summarises_and_alerts_when_accounting_does_not_balance():
    view = build_cluster_report_view({"cluster_report": _report()})
    assert view["kc_summary"]["accounting"].startswith("Not balanced: 2")
    assert "2 of 40 keywords left clustering without a recorded reason" in view["kc_accounting_alert"]
    assert "No cluster was checked against live Google results" in view["kc_verification"]
    assert view["kc_summary"]["size"].startswith("2 small cluster(s) merged; 1 oversized cluster(s) split into 3")
    balanced = build_cluster_report_view({"cluster_report": _report(keyword_accounting={"input_count": 40, "clustered_count": 32, "orphan_count": 8, "balanced": True})})
    assert balanced["kc_accounting_alert"] is None


def test_view_detail_is_capped_and_marks_comparison_and_location():
    view = build_cluster_report_view({"cluster_report": _report()})
    assert len(view["kc_details"]) == DETAIL_CLUSTER_CAP
    assert f"Showing {DETAIL_CLUSTER_CAP} of {DETAIL_CLUSTER_CAP + 3} clusters" in view["kc_details_note"]
    assert view["kc_details"][0]["tone"] == "alert" and view["kc_details"][1]["tone"] is None
    first = view["kc_details"][0]
    assert first["rows"][1][2] == "N/A" and first["rows"][1][1] == "Supporting"
    assert "Why this primary" in first["detail"]
    overview = view["kc_overview"][0]["cells"]
    assert overview[0] == ("Cluster 0", "/blog/kw-0") and overview[3] == ("Comparison", "Consideration Guide")


def test_view_orphans_roadmap_quality_and_provenance():
    view = build_cluster_report_view({"cluster_report": _report()})
    reasons = [r["cells"][3] for r in view["kc_orphans"]]
    assert reasons[1].startswith("Dropped by the model as not relevant: not a service search")
    assert "8 keyword(s) have demand but no cluster" in view["kc_orphans_note"]
    assert [r["cells"][0] for r in view["kc_roadmap"]] == [1, 2]
    issues = [r["cells"][0] for r in view["kc_quality"]]
    assert issues == ["Mixed intent", "Duplicate primary keyword", "No content type", "Single-keyword clusters"]
    assert "2 predict keyword cannibalization" in view["kc_quality_note"]
    assert view["kc_provenance"]["serp_check"] == "Skipped: no resolved market for this client"
    assert view["kc_provenance"]["invented_rejected"] == 3
    assert [r["cells"] for r in view["kc_exclusions"]] == [["A competitor's brand name", 7], ["Mentions a past year", 5]]
    assert view["kc_provenance_notes"] == ["Verified: placed keywords checked", "Inferred: no SERP check"]


def test_view_degrades_to_the_flat_list_for_old_runs():
    view = build_cluster_report_view({"clusters": [{"name": "Old", "primary_keyword": "seo", "keyword_count": 3}]})
    assert view["kc_overview"][0]["cells"][1] == "seo"
    assert view["kc_summary"]["accounting"] == "Not recorded"


# --------------------------------------------------------------- rendering


def _bundle():
    payload = {"card_type": "search_demand_report", "cluster_report": _report(), "summary": "Test run."}
    return {"client": {"name": "Acme"}, "reports": [{"title": "Search Demand", "payload": payload}]}


def test_pdf_renders_every_layer():
    from pypdf import PdfReader

    from app.services.report_pdf import build_reports_pdf

    text = " ".join(p.extract_text() for p in PdfReader(BytesIO(build_reports_pdf(_bundle(), single=True))).pages)
    for heading in ("Keyword Clustering Summary", "Content Hubs", "Cluster Overview", "Cluster Detail",
                    "Orphan Keywords", "Content Roadmap", "Clustering Quality", "Clustering Provenance",
                    "Keyword accounting did not balance", "SEO hub", "Why this primary"):
        assert heading in text, heading
    # The raw report is not dumped as an extra auto-generated section.
    assert "Cluster Report" not in text


def test_docx_renders_every_layer():
    from docx import Document

    from app.services.report_docx import build_reports_docx

    doc = Document(BytesIO(build_reports_docx(_bundle(), single=True)))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    for heading in ("Keyword Clustering Summary", "Content Hubs", "Cluster Detail", "Clustering Quality",
                    "Keyword accounting did not balance", "Cluster 0\n/blog/kw-0"):
        assert heading in text, heading


@pytest.mark.asyncio
async def test_rule_fallback_when_the_model_fails_says_so():
    with patch("app.config.get_settings") as settings, patch(
        "app.services.keyword_llm_clustering.llm_cluster_keywords",
        new=AsyncMock(return_value={"skipped": "mock_or_no_key", "clusters": []}),
    ):
        settings.return_value.keyword_clustering_use_model = True
        report = await run_keyword_clustering([dict(r) for r in ROWS], within_deadline=True)
    prov = report["clustering_provenance"]
    assert prov["grouping"] == "rules" and prov["model_clustering_reason"] == "no model key configured"
    assert report["serp_clustering"]["reason"] == "rule_fallback"
    assert report["keyword_accounting"]["balanced"] is True
