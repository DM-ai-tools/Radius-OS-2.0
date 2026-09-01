"""Word (.docx) report export — same blueprint/normalize layer as the PDF renderer,
so a section fix in report_pdf_document.py should apply to both formats."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from app.services.report_docx import build_reports_docx


def _docx_text(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def test_build_reports_docx_empty_bundle():
    bundle = {
        "client": {"name": "Acme Corp", "primary_url": "https://acme.com"},
        "exported_at": "2026-08-12T12:00:00+00:00",
        "reports": [],
        "phase_statuses": {"overall_readiness_score": 42},
    }
    body = build_reports_docx(bundle)
    assert body[:2] == b"PK"  # docx is a zip archive
    assert len(body) > 500


def test_build_reports_docx_with_report():
    bundle = {
        "client": {"name": "Acme Corp"},
        "exported_at": "2026-08-12T12:00:00+00:00",
        "reports": [
            {
                "title": "Content Strategy: Acme Corp",
                "card_type": "content_strategy_report",
                "payload": {
                    "card_type": "content_strategy_report",
                    "summary": "Three pillars identified.",
                    "priority_queue": [
                        {"title": "Best CRM for startups", "priority": "Quick win", "volume": 1200},
                    ],
                },
            }
        ],
        "phase_statuses": {},
    }
    body = build_reports_docx(bundle, single=True)
    assert body[:2] == b"PK"
    text = _docx_text(body)
    assert "Content Strategy: Acme Corp" in text
    assert "Best CRM for startups" in text


def test_content_audit_docx_shows_themed_inventory():
    bundle = {
        "client": {"name": "Acme"},
        "exported_at": "2026-08-12T12:00:00+00:00",
        "reports": [
            {
                "card_type": "content_audit_report",
                "payload": {
                    "card_type": "content_audit_report",
                    "note": "Locked audit",
                    "inventory": [
                        {"path": "/what-is-seo", "title": "What Is SEO", "theme": "SEO Fundamentals", "disposition": "keep"},
                    ],
                    "themes": [
                        {
                            "theme": "SEO Fundamentals",
                            "url_count": 1,
                            "dispositions": {"keep": 1},
                            "urls": [{"path": "/what-is-seo", "title": "What Is SEO", "disposition": "keep"}],
                        },
                    ],
                },
            }
        ],
        "phase_statuses": {},
    }
    text = _docx_text(build_reports_docx(bundle, single=True))
    assert "Themed Inventory" in text
    assert "SEO Fundamentals" in text
    assert "keep: 1" in text


def test_content_planning_docx_shows_funnel_and_content_signals_no_duplicate_roadmap():
    bundle = {
        "client": {"name": "Acme"},
        "exported_at": "2026-08-12T12:00:00+00:00",
        "reports": [
            {
                "card_type": "content_planning_report",
                "payload": {
                    "card_type": "content_planning_report",
                    "note": "Locked roadmap",
                    "pages": [
                        {
                            "priority_rank": 1,
                            "url_n": "/seo-pricing",
                            "primary_keyword": "seo pricing",
                            "action": "create",
                            "priority_tier": "quick_win",
                            "funnel": "BOFU",
                            "angle": "comparison",
                            "business_fit": {"score": 90, "reason": "core_service"},
                            "competitor_domains": ["rival-a.com"],
                        }
                    ],
                    "roadmap": [
                        {
                            "priority_rank": 1,
                            "url_n": "/seo-pricing",
                            "primary_keyword": "seo pricing",
                            "action": "create",
                            "priority_tier": "quick_win",
                            "funnel": "BOFU",
                        }
                    ],
                    "excluded": [],
                },
            }
        ],
        "phase_statuses": {},
    }
    text = _docx_text(build_reports_docx(bundle, single=True))
    assert "Funnel" in text
    assert "BOFU" in text
    assert "Content Signals" in text
    assert "comparison" in text
    assert "rival-a.com" in text
    assert "Merged pages" in text
    assert "Roadmap" not in text


def test_content_production_docx_shows_briefs_and_funnel_balance_once():
    bundle = {
        "client": {"name": "Acme"},
        "exported_at": "2026-08-12T12:00:00+00:00",
        "reports": [
            {
                "card_type": "content_production_report",
                "payload": {
                    "card_type": "content_production_report",
                    "note": "Briefed queue",
                    "briefs": [
                        {
                            "keyword": "seo pricing",
                            "url": "/seo-pricing",
                            "funnel": "BOFU",
                            "angle": "comparison",
                            "status": "ready",
                        }
                    ],
                    "drafts": [],
                    "funnel_balance": {"TOFU": 0, "MOFU": 0, "BOFU": 1},
                    "funnel_balance_warnings": ["No TOFU topics in this plan — funnel is unbalanced."],
                },
            }
        ],
        "phase_statuses": {},
    }
    text = _docx_text(build_reports_docx(bundle, single=True))
    assert "Briefed Topics" in text
    assert "seo pricing" in text
    assert "comparison" in text
    assert "Funnel Balance Warnings" in text
    assert "3. Funnel Balance" in text


def _url_map_fixture_payload() -> dict:
    return {
        "card_type": "site_architecture_blueprint",
        "title": "Site Architecture Blueprint: Acme Corp",
        "client": "Acme Corp",
        "primary_url": "https://acme.com",
        "executive_summary": "Blueprint with URL mapping.",
        "current_state": {"urls_crawled": 50, "status_200": 48},
        "url_map_report": {
            "final_url_map": [
                {
                    "level": 2,
                    "l1_category": "Services",
                    "l2_subcategory": "Local Seo Melbourne",
                    "l3_subsubcategory": None,
                    "l4_attribution": None,
                    "current_url": "/services/local-seo-melbourne",
                    "proposed_url": None,
                    "status": "Optimize Existing",
                    "primary_keyword": "local seo melbourne",
                    "search_volume": 1200,
                    "cpc": 4.5,
                    "secondary_keywords_sheet": [
                        "seo agency melbourne",
                        "melbourne seo services",
                    ],
                    "combined_cluster_volume": 2000,
                    "page_type": "service",
                    "priority": "High",
                    "notes": "SERP validated (service)",
                    "est_products": 2,
                    "in_scope": "Yes",
                }
            ],
        },
    }


def test_site_architecture_docx_shows_url_mapping_columns():
    bundle = {
        "client": {"name": "Acme Corp"},
        "exported_at": "2026-08-12T12:00:00+00:00",
        "reports": [
            {
                "card_type": "site_architecture_blueprint",
                "payload": _url_map_fixture_payload(),
            }
        ],
        "phase_statuses": {},
    }
    text = _docx_text(build_reports_docx(bundle, single=True))
    assert "URL Mapping — Taxonomy & Status" in text
    assert "URL Mapping — Keywords & Opportunity" in text
    assert "local seo melbourne" in text
    assert "Combined Vol" in text
    assert "In Scope?" in text
    assert "Secondary Keywords (5-10)" in text
    assert "Optimize Existing" in text


# --- API dispatch: format query param must select the right builder ----------------

def test_report_export_formats_dispatch():
    from app.api.clients import REPORT_EXPORT_FORMATS, _build_export_body

    assert set(REPORT_EXPORT_FORMATS) == {"pdf", "docx"}
    bundle = {
        "client": {"name": "Acme"},
        "exported_at": "2026-08-12T12:00:00+00:00",
        "reports": [],
        "phase_statuses": {},
    }
    pdf_body = _build_export_body(bundle, fmt="pdf", single=False)
    assert pdf_body[:4] == b"%PDF"
    docx_body = _build_export_body(bundle, fmt="docx", single=False)
    assert docx_body[:2] == b"PK"
