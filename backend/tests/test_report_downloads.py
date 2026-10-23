"""Downloads: redesigned PDF/Word reports and the keyword research CSV.

PDF: a drawn cover (title, client, date, contents) and, on every later page,
a running header and "Page X of Y". Word: same cover and frame, with every
property element in schema order so Word opens the file. CSV: one row per
keyword, offered for keyword research only, safe to open in a spreadsheet.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest
from fastapi import HTTPException

from app.services.report_csv import COLUMNS, build_keyword_research_csv


def _search_demand_payload():
    return {
        "card_type": "search_demand_report",
        "summary": "Keyword research for Acme.",
        "cluster_report": {
            "clusters": [
                {
                    "name": "SEO Agency", "primary_keyword": "seo agency", "recommended_url": "/blog/seo-agency",
                    "pillar_name": "SEO", "modifier_label": "Core topic", "keyword_count": 2,
                    "keywords": [
                        {"keyword": "seo agency", "role": "Primary", "volume": 1300, "difficulty": 40,
                         "intent": "commercial", "intent_source": "provider", "intent_confidence": "high",
                         "funnel": "BOFU", "funnel_source": "heuristic", "opportunity_score": 70},
                        {"keyword": "=cmd|calc", "role": "Secondary", "volume": None, "difficulty": None,
                         "intent": "commercial", "funnel": "BOFU"},
                    ],
                }
            ],
            "orphans": [{"keyword": "seo podcast", "volume": 30, "intent": "informational",
                         "exclusion_reason": "dropped_for_relevance", "notes": "not a service"}],
        },
        "keyword_dataset": [
            {"keyword": "seo agency", "service": "SEO", "seed": "seo", "cpc": 12.5},
            {"keyword": "seo tips", "volume": 700, "volume_estimated": True, "service": "SEO", "intent": "informational"},
        ],
    }


# --------------------------------------------------------------- CSV


def _parse(body: bytes) -> list[dict[str, str]]:
    assert body.startswith(b"\xef\xbb\xbf"), "BOM so Excel reads UTF-8"
    return list(csv.DictReader(io.StringIO(body.decode("utf-8-sig"))))


def test_csv_has_one_row_per_keyword_with_status_and_labels():
    rows = _parse(build_keyword_research_csv(_search_demand_payload()))
    assert list(rows[0].keys()) == COLUMNS
    by_kw = {r["Keyword"]: r for r in rows}
    agency = by_kw["seo agency"]
    assert (agency["Status"], agency["Cluster"], agency["Role"], agency["Content hub"]) == (
        "Clustered", "SEO Agency", "Primary", "SEO")
    assert (agency["Search volume"], agency["Volume status"], agency["CPC"], agency["Service"]) == (
        "1300", "Measured", "12.5", "SEO")
    assert (agency["Intent"], agency["Intent source"], agency["Funnel stage"]) == ("Commercial", "provider", "BOFU")
    orphan = by_kw["seo podcast"]
    assert orphan["Status"] == "Orphan" and orphan["Why it has no cluster"].startswith("Dropped by the model")
    extra = by_kw["seo tips"]
    assert extra["Status"] == "Not in cluster map" and extra["Volume status"] == "Estimated"
    assert len(rows) == 4


def test_csv_marks_missing_values_and_blocks_formulas():
    rows = _parse(build_keyword_research_csv(_search_demand_payload()))
    risky = next(r for r in rows if r["Keyword"].endswith("cmd|calc"))
    assert risky["Keyword"] == "'=cmd|calc"
    assert (risky["Search volume"], risky["Volume status"], risky["Keyword difficulty"]) == ("N/A", "Not measured", "N/A")


# --------------------------------------------------------------- endpoint


def _bundle(*payloads):
    return {
        "client": {"name": "Acme", "primary_url": "acme.com.au"},
        "exported_at": "2026-10-19T00:00:00Z",
        "filename_slug": "acme",
        "reports": [{"card_type": p["card_type"], "title": p["card_type"], "payload": p} for p in payloads],
    }


@pytest.mark.asyncio
async def test_csv_is_served_for_keyword_research_and_refused_elsewhere():
    from app.api.clients import export_client_reports, export_single_client_report

    bundle = _bundle(_search_demand_payload(), {"card_type": "technical_seo_report", "summary": "x"})
    with patch("app.services.report_export.build_report_export", new=AsyncMock(return_value=bundle)):
        response = await export_single_client_report(
            client_id="00000000-0000-0000-0000-000000000001", card_type="search_demand_report",
            format="csv", db=None, _=None)
        assert response.media_type.startswith("text/csv")
        assert re.search(r'filename="acme-.*\.csv"', response.headers["content-disposition"])
        assert b"seo agency" in response.body
        with pytest.raises(HTTPException) as other:
            await export_single_client_report(
                client_id="00000000-0000-0000-0000-000000000001", card_type="technical_seo_report",
                format="csv", db=None, _=None)
        assert other.value.status_code == 400
        with pytest.raises(HTTPException):
            await export_client_reports(client_id="00000000-0000-0000-0000-000000000001", format="csv", db=None, _=None)


# --------------------------------------------------------------- PDF


def test_pdf_cover_contents_and_page_frame():
    from pypdf import PdfReader

    from app.services.report_pdf import build_reports_pdf

    pdf = PdfReader(BytesIO(build_reports_pdf(_bundle(_search_demand_payload()), single=True)))
    cover = pdf.pages[0].extract_text()
    for text in ("Radius OS", "search_demand_report", "Acme", "In this report", "19 October 2026"):
        assert text in cover, text
    assert "Keyword Clustering Summary" in cover  # contents list names the sections
    total = len(pdf.pages)
    assert "Page" not in cover  # the cover carries no page number
    second = pdf.pages[1].extract_text()
    assert f"Page 2 of {total}" in second and "acme.com.au" in second
    assert "Summary" in second


def test_pdf_bundle_lists_the_reports_on_the_cover():
    from pypdf import PdfReader

    from app.services.report_pdf import build_reports_pdf

    bundle = _bundle(_search_demand_payload(), {"card_type": "technical_seo_report", "summary": "Tech."})
    pdf = PdfReader(BytesIO(build_reports_pdf(bundle, single=False)))
    cover = pdf.pages[0].extract_text()
    assert "2 phase reports in this pack" in cover
    assert "1. search_demand_report" in cover and "2. technical_seo_report" in cover


# --------------------------------------------------------------- Word


def test_docx_cover_frame_and_schema_order():
    from docx import Document
    from lxml import etree

    from app.services.report_docx import _ORDER, build_reports_docx

    body = build_reports_docx(_bundle(_search_demand_payload()), single=True)
    doc = Document(BytesIO(body))
    assert len(doc.sections) == 2
    assert not "".join(p.text for p in doc.sections[0].header.paragraphs).strip()
    assert doc.sections[1].header.paragraphs[0].text == "search_demand_report\tacme.com.au"
    cover = doc.tables[0].rows[0].cells[0].text
    assert "Radius OS" in cover and "Acme" in cover
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "In this report" in text and "Keyword Clustering Summary" in text

    z = zipfile.ZipFile(BytesIO(body))
    footer = next(z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))
    assert " PAGE " in footer and " NUMPAGES " in footer
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for name in [n for n in z.namelist() if re.match(r"word/(document|header\d*|footer\d*)\.xml", n)]:
        root = etree.fromstring(z.read(name))
        for tag, order in _ORDER.items():
            for el in root.iter(ns + tag):
                ranks = [order.index(c.tag.split("}")[-1]) for c in el if c.tag.split("}")[-1] in order]
                assert ranks == sorted(ranks), (name, tag)
